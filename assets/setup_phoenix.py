#!/usr/bin/env python3
"""
PhoenixOffice Core & Leitstand-Frontend Setup Script
Erzeugt das komplette Backend, Frontend-Templates, statische Assets
und OfficeWrite-Pipelines in einem Durchlauf.
"""

from pathlib import Path

FILES = {}

# -------------------------------------------------------------
# Python Package Markers (__init__.py)
# -------------------------------------------------------------
PACKAGE_DIRS = [
    "src",
    "src/app",
    "src/db",
    "src/ai_bridge",
    "src/officewrite",
    "tests"
]
for d in PACKAGE_DIRS:
    FILES[f"{d}/__init__.py"] = '"""Package init."""\n'

# -------------------------------------------------------------
# Dependencies & Config
# -------------------------------------------------------------
FILES["requirements.txt"] = """fastapi==0.115.0
uvicorn[standard]==0.30.0
sqlalchemy==2.0.35
pydantic==2.9.2
python-docx==1.1.2
jinja2==3.1.4
aiofiles==24.1.0
requests==2.32.3
pytest==8.3.3
email-validator==2.2.0
"""

FILES[".env.example"] = """APP_NAME="PhoenixOffice Leitstand"
ENV="development"
PORT=8000
DATABASE_URL="sqlite:///./phoenixoffice.db"
STORAGE_DIR="./storage_vault"
DS_AI_ENDPOINT="http://localhost:11434/api/generate"
DS_AI_ENABLED="false"
"""

FILES[".gitignore"] = """__pycache__/
*.pyc
*.db
.env
storage_vault/
.pytest_cache/
.venv/
"""

# -------------------------------------------------------------
# DB Models & Session
# -------------------------------------------------------------
FILES["src/db/models.py"] = '''import uuid
from datetime import datetime
from sqlalchemy import Column, String, DateTime, Text, ForeignKey, JSON
from sqlalchemy.orm import declarative_base, relationship

Base = declarative_base()

class EmailRecord(Base):
    __tablename__ = "emails"

    id = Column(String(36), primary_key=True, default=lambda: str(uuid.uuid4()))
    message_id = Column(String(255), unique=True, index=True, nullable=True)
    sender = Column(String(255), nullable=False)
    recipient = Column(String(255), nullable=False)
    subject = Column(String(500), nullable=False)
    body_text = Column(Text, nullable=False)
    received_at = Column(DateTime, default=datetime.utcnow)
    status = Column(String(50), default="ingested")
    intent_tag = Column(String(100), nullable=True)

    documents = relationship("DocumentRecord", back_populates="email")

class DocumentRecord(Base):
    __tablename__ = "documents"

    id = Column(String(36), primary_key=True, default=lambda: str(uuid.uuid4()))
    source_email_id = Column(String(36), ForeignKey("emails.id"), nullable=True)
    title = Column(String(255), nullable=False)
    doc_type = Column(String(100), default="generic")
    storage_path = Column(String(500), nullable=False)
    created_at = Column(DateTime, default=datetime.utcnow)
    metadata_payload = Column(JSON, nullable=True)

    email = relationship("EmailRecord", back_populates="documents")
'''

FILES["src/db/session.py"] = '''import os
from sqlalchemy import create_engine
from sqlalchemy.orm import sessionmaker
from src.db.models import Base

DB_URL = os.getenv("DATABASE_URL", "sqlite:///./phoenixoffice.db")

engine = create_engine(
    DB_URL,
    connect_args={"check_same_thread": False} if DB_URL.startswith("sqlite") else {}
)
SessionLocal = sessionmaker(autocommit=False, autoflush=False, bind=engine)

def init_db():
    Base.metadata.create_all(bind=engine)

def get_db():
    db = SessionLocal()
    try:
        yield db
    finally:
        db.close()
'''

# -------------------------------------------------------------
# AI Bridge
# -------------------------------------------------------------
FILES["src/ai_bridge/classifier.py"] = '''import os
import requests

AI_ENDPOINT = os.getenv("DS_AI_ENDPOINT", "http://localhost:11434/api/generate")
AI_ENABLED = os.getenv("DS_AI_ENABLED", "false").lower() == "true"

def extract_intent_and_summary(subject: str, body: str) -> dict:
    if AI_ENABLED:
        try:
            prompt = (
                f"Klassifiziere den Vorgang und fasse zusammen.\\n"
                f"Betreff: {subject}\\nText: {body[:1500]}\\n"
                f"Format: INTENT: [vertrag|briefing|lagebericht|allgemein]; SUMMARY: [kurztext]"
            )
            resp = requests.post(
                AI_ENDPOINT,
                json={"model": "mistral", "prompt": prompt, "stream": False},
                timeout=3.0
            )
            if resp.status_code == 200:
                raw_text = resp.json().get("response", "")
                return {"intent": "ai_classified", "summary": raw_text.strip(), "source": "linux_mint_ai"}
        except Exception:
            pass

    combined = (subject + " " + body).lower()
    if any(k in combined for k in ["lage", "lagebericht", "sicherheit", "einsatz"]):
        intent = "lagebericht"
    elif any(k in combined for k in ["vertrag", "nda", "vereinbarung", "unterschrift"]):
        intent = "vertrag"
    elif any(k in combined for k in ["briefing", "konzept", "anforderung", "architektur"]):
        intent = "briefing"
    else:
        intent = "allgemeine_anfrage"

    first_line = body.strip().split("\\n")[0][:180] if body.strip() else subject
    return {"intent": intent, "summary": first_line, "source": "rule_heuristic"}
'''

# -------------------------------------------------------------
# OfficeWrite Module
# -------------------------------------------------------------
FILES["src/officewrite/generator.py"] = '''from pathlib import Path
from datetime import datetime
from docx import Document
from docx.shared import Pt, Inches, RGBColor

class OfficeWriteEngine:
    def __init__(self, storage_vault: str = "./storage_vault"):
        self.storage_vault = Path(storage_vault)
        self.storage_vault.mkdir(parents=True, exist_ok=True)

    def build_document(self, doc_id: str, title: str, doc_type: str, author: str, summary: str, body: str) -> str:
        doc = Document()
        style = doc.styles['Normal']
        font = style.font
        font.name = 'Calibri'
        font.size = Pt(11)

        h1 = doc.add_heading(f"PhoenixOffice :: {doc_type.upper()}", level=0)
        h1.style.font.color.rgb = RGBColor(0x00, 0x33, 0x66)

        doc.add_paragraph(f"Dokumenten-ID: {doc_id}")
        doc.add_paragraph(f"Erstellt am: {datetime.utcnow().strftime('%Y-%m-%d %H:%M:%S')} UTC")
        doc.add_paragraph(f"Autor/Modul: {author}")
        doc.add_paragraph(f"Referenz-Titel: {title}")
        doc.add_paragraph("-" * 50)

        doc.add_heading("1. Executive Summary & Einstufung", level=1)
        p_summary = doc.add_paragraph(summary)
        p_summary.paragraph_format.left_indent = Inches(0.2)

        doc.add_heading("2. E-Mail Quelltext / Payload", level=1)
        doc.add_paragraph(body)

        target_path = self.storage_vault / f"{doc_id}_{doc_type}.docx"
        doc.save(str(target_path))
        return str(target_path)
'''

# -------------------------------------------------------------
# Frontend Static Assets & HTML Templates
# -------------------------------------------------------------
FILES["src/app/static/styles.css"] = """:root {
  --bg: #0b0f19;
  --surface: #111827;
  --border: #1f2937;
  --primary: #0284c7;
  --primary-hover: #0369a1;
  --text: #f3f4f6;
  --muted: #9ca3af;
  --success: #10b981;
}
* { box-sizing: border-box; margin: 0; padding: 0; }
body {
  font-family: -apple-system, BlinkMacSystemFont, "Segoe UI", Roboto, sans-serif;
  background-color: var(--bg);
  color: var(--text);
  line-height: 1.5;
}
.header {
  background: var(--surface);
  border-bottom: 1px solid var(--border);
  padding: 1rem 2rem;
  display: flex;
  justify-content: space-between;
  align-items: center;
}
.nav a {
  color: var(--muted);
  text-decoration: none;
  margin-left: 1.5rem;
  font-weight: 500;
}
.nav a:hover, .nav a.active { color: var(--primary); }
.container { max-width: 1000px; margin: 2rem auto; padding: 0 1rem; }
.card {
  background: var(--surface);
  border: 1px solid var(--border);
  border-radius: 8px;
  padding: 1.5rem;
  margin-bottom: 1.5rem;
}
h1, h2 { margin-bottom: 1rem; }
.badge {
  background: rgba(16, 185, 129, 0.1);
  color: var(--success);
  border: 1px solid var(--success);
  padding: 0.2rem 0.6rem;
  border-radius: 9999px;
  font-size: 0.8rem;
}
.form-group { margin-bottom: 1rem; }
label { display: block; font-size: 0.85rem; color: var(--muted); margin-bottom: 0.3rem; }
input, textarea {
  width: 100%;
  padding: 0.6rem;
  background: #1f2937;
  border: 1px solid var(--border);
  border-radius: 4px;
  color: #fff;
  font-size: 0.95rem;
}
button {
  background: var(--primary);
  color: #fff;
  border: none;
  padding: 0.6rem 1.2rem;
  border-radius: 4px;
  cursor: pointer;
  font-weight: 600;
}
button:hover { background: var(--primary-hover); }
table { width: 100%; border-collapse: collapse; margin-top: 1rem; }
th, td { text-align: left; padding: 0.75rem; border-bottom: 1px solid var(--border); font-size: 0.9rem; }
th { color: var(--muted); }
"""

FILES["src/app/templates/index.html"] = """<!DOCTYPE html>
<html lang="de">
<head>
  <meta charset="utf-8">
  <title>PhoenixOffice — Leitstand</title>
  <link rel="stylesheet" href="/static/styles.css">
</head>
<body>
  <header class="header">
    <h2>PhoenixOffice :: Control Center</h2>
    <nav class="nav">
      <a href="/" class="active">Dashboard</a>
      <a href="/ingest">E-Mail Ingest</a>
    </nav>
  </header>
  <main class="container">
    <div class="card">
      <h2>Systemstatus</h2>
      <p>Subsystem-Health: <span class="badge">Online</span></p>
      <div style="margin-top: 1rem;">
        <p><strong>DB-Status:</strong> Verbunden (SQLite)</p>
        <p><strong>OfficeWrite Engine:</strong> Bereit (DOCX Builder v1.0)</p>
        <p><strong>Data Science Bridge:</strong> Heuristik aktiv (Ollama Fallback)</p>
      </div>
    </div>
    <div class="card">
      <h2>Neueste Vorgänge</h2>
      <table>
        <thead>
          <tr>
            <th>ID</th>
            <th>Absender</th>
            <th>Betreff</th>
            <th>Status</th>
            <th>Aktion</th>
          </tr>
        </thead>
        <tbody>
          {% for email in emails %}
          <tr>
            <td>{{ email.id[:8] }}...</td>
            <td>{{ email.sender }}</td>
            <td>{{ email.subject }}</td>
            <td><span class="badge">{{ email.status }}</span></td>
            <td><a href="/status/{{ email.id }}" style="color: var(--primary);">Details</a></td>
          </tr>
          {% else %}
          <tr>
            <td colspan="5" style="color: var(--muted); text-align: center;">Keine Vorgänge im Ingestion-Puffer.</td>
          </tr>
          {% endfor %}
        </tbody>
      </table>
    </div>
  </main>
</body>
</html>
"""

FILES["src/app/templates/ingest.html"] = """<!DOCTYPE html>
<html lang="de">
<head>
  <meta charset="utf-8">
  <title>Ingest — PhoenixOffice</title>
  <link rel="stylesheet" href="/static/styles.css">
</head>
<body>
  <header class="header">
    <h2>PhoenixOffice :: Control Center</h2>
    <nav class="nav">
      <a href="/">Dashboard</a>
      <a href="/ingest" class="active">E-Mail Ingest</a>
    </nav>
  </header>
  <main class="container">
    <div class="card">
      <h2>E-Mail erfassen & OfficeWrite anstoßen</h2>
      <form action="/ingest" method="post">
        <div class="form-group">
          <label>Absender (From)</label>
          <input type="email" name="sender" required value="kontakt@philosoph.online">
        </div>
        <div class="form-group">
          <label>Empfänger (To)</label>
          <input type="text" name="recipient" required value="leitstand@phoenixoffice.org">
        </div>
        <div class="form-group">
          <label>Betreff (Subject)</label>
          <input type="text" name="subject" required value="Lagebericht: Infrastruktur & Schnittstellen">
        </div>
        <div class="form-group">
          <label>Nachrichtentext (Body)</label>
          <textarea name="body_text" rows="6" required>Priorität Inbound verifiziert. Vorgang zur automatischen Ablage und OfficeWrite-Generierung übergeben.</textarea>
        </div>
        <button type="submit">E-Mail verarbeiten</button>
      </form>
    </div>
  </main>
</body>
</html>
"""

FILES["src/app/templates/status.html"] = """<!DOCTYPE html>
<html lang="de">
<head>
  <meta charset="utf-8">
  <title>Vorgangsstatus — PhoenixOffice</title>
  <link rel="stylesheet" href="/static/styles.css">
</head>
<body>
  <header class="header">
    <h2>PhoenixOffice :: Vorgang {{ email.id[:8] }}</h2>
    <nav class="nav">
      <a href="/">Zurück zur Übersicht</a>
    </nav>
  </header>
  <main class="container">
    <div class="card">
      <h2>Metadaten</h2>
      <p><strong>Absender:</strong> {{ email.sender }}</p>
      <p><strong>Betreff:</strong> {{ email.subject }}</p>
      <p><strong>Einstufung / Intent:</strong> <span class="badge">{{ email.intent_tag or 'in Bearbeitung' }}</span></p>
      <p><strong>Status:</strong> {{ email.status }}</p>
    </div>
    <div class="card">
      <h2>OfficeWrite Dokumente</h2>
      {% for doc in email.documents %}
      <div style="padding: 0.5rem 0; border-bottom: 1px solid var(--border);">
        <p><strong>Titel:</strong> {{ doc.title }}</p>
        <p><strong>Ablagepfad:</strong> <code>{{ doc.storage_path }}</code></p>
      </div>
      {% else %}
      <p style="color: var(--muted);">Dokumenterzeugung läuft im Hintergrund...</p>
      {% endfor %}
    </div>
  </main>
</body>
</html>
"""

# -------------------------------------------------------------
# FastAPI App (API + Web Frontend)
# -------------------------------------------------------------
FILES["src/app/main.py"] = '''import uuid
from typing import Optional
from pathlib import Path
from fastapi import FastAPI, Depends, HTTPException, BackgroundTasks, Request, Form
from fastapi.responses import HTMLResponse, RedirectResponse
from fastapi.staticfiles import StaticFiles
from fastapi.templating import Jinja2Templates
from pydantic import BaseModel, EmailStr
from sqlalchemy.orm import Session

from src.db.session import init_db, get_db
from src.db.models import EmailRecord, DocumentRecord
from src.ai_bridge.classifier import extract_intent_and_summary
from src.officewrite.generator import OfficeWriteEngine

init_db()
app = FastAPI(title="PhoenixOffice Core & Leitstand", version="1.0.0")

BASE_DIR = Path(__file__).resolve().parent
app.mount("/static", StaticFiles(directory=str(BASE_DIR / "static")), name="static")
templates = Jinja2Templates(directory=str(BASE_DIR / "templates"))

write_engine = OfficeWriteEngine()

class EmailIngestSchema(BaseModel):
    sender: EmailStr
    recipient: str
    subject: str
    body_text: str
    message_id: Optional[str] = None

def process_email_background(email_id: str):
    from src.db.session import SessionLocal
    db = SessionLocal()
    try:
        rec = db.query(EmailRecord).filter(EmailRecord.id == email_id).first()
        if not rec:
            return

        analysis = extract_intent_and_summary(rec.subject, rec.body_text)
        rec.intent_tag = analysis["intent"]
        rec.status = "processed"

        doc_id = str(uuid.uuid4())
        doc_path = write_engine.build_document(
            doc_id=doc_id,
            title=rec.subject,
            doc_type=rec.intent_tag,
            author="PhoenixOffice Ingest Engine",
            summary=analysis["summary"],
            body=rec.body_text
        )

        doc_entry = DocumentRecord(
            id=doc_id,
            source_email_id=rec.id,
            title=f"{rec.intent_tag.upper()}: {rec.subject}",
            doc_type=rec.intent_tag,
            storage_path=doc_path,
            metadata_payload=analysis
        )
        db.add(doc_entry)
        db.commit()
    finally:
        db.close()

# --- Web Frontend Routes ---
@app.get("/", response_class=HTMLResponse)
def page_dashboard(request: Request, db: Session = Depends(get_db)):
    emails = db.query(EmailRecord).order_by(EmailRecord.received_at.desc()).limit(15).all()
    return templates.TemplateResponse("index.html", {"request": request, "emails": emails})

@app.get("/ingest", response_class=HTMLResponse)
def page_ingest_form(request: Request):
    return templates.TemplateResponse("ingest.html", {"request": request})

@app.post("/ingest")
def handle_web_ingest(
    bg: BackgroundTasks,
    sender: str = Form(...),
    recipient: str = Form(...),
    subject: str = Form(...),
    body_text: str = Form(...),
    db: Session = Depends(get_db)
):
    record_id = str(uuid.uuid4())
    db_email = EmailRecord(
        id=record_id,
        message_id=f"web-{record_id[:8]}",
        sender=sender,
        recipient=recipient,
        subject=subject,
        body_text=body_text,
        status="queued"
    )
    db.add(db_email)
    db.commit()
    bg.add_task(process_email_background, record_id)
    return RedirectResponse(url=f"/status/{record_id}", status_code=303)

@app.get("/status/{email_id}", response_class=HTMLResponse)
def page_status(email_id: str, request: Request, db: Session = Depends(get_db)):
    email = db.query(EmailRecord).filter(EmailRecord.id == email_id).first()
    if not email:
        raise HTTPException(status_code=404, detail="Vorgang nicht gefunden")
    return templates.TemplateResponse("status.html", {"request": request, "email": email})

# --- JSON API Endpoints ---
@app.get("/api/health")
def health():
    return {
        "status": "ok",
        "subsystems": {
            "officewrite": "online",
            "db": "ready",
            "ai_bridge": "connected"
        }
    }

@app.post("/api/emails/ingest", status_code=202)
def api_ingest(payload: EmailIngestSchema, bg: BackgroundTasks, db: Session = Depends(get_db)):
    record_id = str(uuid.uuid4())
    db_email = EmailRecord(
        id=record_id,
        message_id=payload.message_id or f"api-{record_id[:8]}",
        sender=payload.sender,
        recipient=payload.recipient,
        subject=payload.subject,
        body_text=payload.body_text,
        status="queued"
    )
    db.add(db_email)
    db.commit()
    bg.add_task(process_email_background, record_id)
    return {"status": "accepted", "email_id": record_id}
'''

# -------------------------------------------------------------
# Automated Smoke Test
# -------------------------------------------------------------
FILES["tests/test_flow.py"] = '''import time
from fastapi.testclient import TestClient
from src.app.main import app

client = TestClient(app)

def test_full_pipeline():
    res = client.get("/api/health")
    assert res.status_code == 200

    html = client.get("/")
    assert html.status_code == 200
    assert "PhoenixOffice :: Control Center" in html.text

    payload = {
        "sender": "lagezentrum@bundeswehr-partner.de",
        "recipient": "ops@phoenixoffice.org",
        "subject": "Lagebericht: Test",
        "body_text": "Infrastrukturanalyse Durchlauf."
    }
    ingest_res = client.post("/api/emails/ingest", json=payload)
    assert ingest_res.status_code == 202
'''

def main():
    root_dir = Path(".")
    print(f"[*] Initialisiere Dateistruktur in: {root_dir.resolve()}")
    for filepath, content in FILES.items():
        full_path = root_dir / filepath
        full_path.parent.mkdir(parents=True, exist_ok=True)
        with open(full_path, "w", encoding="utf-8") as f:
            f.write(content)
        print(f"  -> Angelegt: {filepath}")
    print("\n[✔] Setup abgeschlossen. Alle Dateien, Templates und Assets erzeugt.")

if __name__ == "__main__":
    main()
